# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_whale_brief_doc():
    doc = docx.Document()

    # Set Margins (2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
        # Configure Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run("Brief website Whale Interior — Trang ")
        f_run.font.name = "Arial"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(120, 120, 120)
        
        # Add Page field in Word XML
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        f_p._p.append(fldSimple)

    # Style definitions
    NAVY = RGBColor(10, 25, 47)      # #0A192F
    RED_TITLE = RGBColor(185, 28, 28) # #B91C1C
    GOLD = RGBColor(197, 160, 89)    # #C5A059
    DARK_GRAY = RGBColor(30, 41, 59)
    MUTED_GRAY = RGBColor(100, 116, 139)
    BLUE_LINK = RGBColor(29, 78, 216)

    def set_cell_background(cell, hex_color):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def style_table(table, col_widths, header_bg="B91C1C"):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            for j, cell in enumerate(row.cells):
                cell.width = col_widths[j]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                
                if i == 0:
                    set_cell_background(cell, header_bg)
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for r in p.runs:
                            r.font.bold = True
                            r.font.name = "Arial"
                            r.font.size = Pt(10)
                            r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.name = "Arial"
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = DARK_GRAY

    # Document Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(2)
    r1 = p_title.add_run("BRIEF WEBSITE\n")
    r1.font.bold = True
    r1.font.size = Pt(18)
    r1.font.name = "Arial"
    r1.font.color.rgb = RED_TITLE

    r2 = p_title.add_run("WHALE INTERIOR – NHA TRANG")
    r2.font.bold = True
    r2.font.size = Pt(16)
    r2.font.name = "Arial"
    r2.font.color.rgb = NAVY

    # Demo Links Block
    p_link = doc.add_paragraph()
    p_link.paragraph_format.space_before = Pt(8)
    p_link.paragraph_format.space_after = Pt(4)
    r_l1 = p_link.add_run("Link demo Homepage: ")
    r_l1.font.bold = True
    r_l1.font.name = "Arial"
    r_l1.font.size = Pt(10.5)
    r_l1.font.color.rgb = BLUE_LINK
    r_l2 = p_link.add_run("https://thien230314-bit.github.io/whale-interior-demo/")
    r_l2.font.name = "Arial"
    r_l2.font.size = Pt(10.5)
    r_l2.font.color.rgb = BLUE_LINK
    r_l2.font.underline = True

    p_git = doc.add_paragraph()
    p_git.paragraph_format.space_after = Pt(18)
    r_g1 = p_git.add_run("Mã nguồn (GitHub): ")
    r_g1.font.bold = True
    r_g1.font.name = "Arial"
    r_g1.font.size = Pt(10.5)
    r_g1.font.color.rgb = BLUE_LINK
    r_g2 = p_git.add_run("https://github.com/thien230314-bit/whale-interior-demo")
    r_g2.font.name = "Arial"
    r_g2.font.size = Pt(10.5)
    r_g2.font.color.rgb = BLUE_LINK
    r_g2.font.underline = True

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.bold = True
        run.font.size = Pt(13.5)
        run.font.name = "Arial"
        run.font.color.rgb = RED_TITLE
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.bold = True
        run.font.size = Pt(11.5)
        run.font.name = "Arial"
        run.font.color.rgb = NAVY
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_GRAY
        return p

    # --- SECTION 1 ---
    add_h1("1. Giả định khách hàng")
    add_bullet("Theo yêu cầu của bài tập, chọn một ngành trong các ngành trọng điểm: bất động sản, spa, nhà hàng, du lịch lữ hành hoặc thiết kế kiến trúc – nội thất tại Nha Trang. Ngành được lựa chọn là Thiết kế & Thi công nội thất.")
    add_bullet("Tình huống giả định: một kiến trúc sư kiêm chủ xưởng sản xuất nội thất quy mô vừa tại Nha Trang tìm đến người thiết kế web và nhờ làm một website thương hiệu cao cấp cho công ty. Toàn bộ thông tin về chủ doanh nghiệp và thương hiệu bên dưới là giả định để phục vụ bài tập.")

    add_h2("1.1. Hồ sơ khách hàng giả định")
    t1 = doc.add_table(rows=9, cols=2)
    data_t1 = [
        ["Thông tin", "Nội dung giả định"],
        ["Ngành", "Thiết kế kiến trúc & Thi công nội thất cao cấp (Villa, Penthouse, Nhà phố)"],
        ["Tên thương hiệu", "WHALE INTERIOR (Sở hữu xưởng chế tác nội thất Whale Craft)"],
        ["Người nhờ làm web", "Anh Minh, 38 tuổi, Kiến trúc sư trưởng kiêm Giám đốc điều hành"],
        ["Địa điểm", "Nha Trang: Showroom tại đường Trần Phú (TP. Nha Trang); Xưởng sản xuất 1.500m² tại Cụm Công nghiệp Diên Phú (Diên Khánh, Khánh Hòa); Chi nhánh VP tại TP.HCM"],
        ["Quy mô", "Khoảng 35 nhân sự (Kiến trúc sư, kỹ sư giám sát công trường, thợ mộc mộc máy CNC và đội thợ sơn PU hoàn thiện)"],
        ["Cách tiếp cận khách hàng hiện nay", "Chủ yếu qua khách hàng cũ giới thiệu truyền miệng, Fanpage Facebook, Zalo cá nhân; gửi file PDF hồ sơ năng lực 50MB qua tin nhắn để khách xem"],
        ["Hiểu biết về web", "Rất ít; thành thạo phần mềm thiết kế chuyên ngành (AutoCAD, 3ds Max, SketchUp), nhưng không biết gì về kỹ thuật lập trình web"],
        ["Mong muốn về thời gian", "Có website chạy chính thức trong khoảng 1–2 tháng, kịp trước mùa cao điểm hoàn thiện nhà cuối năm và mùa du lịch đón đầu dự án nghỉ dưỡng"]
    ]
    for r_idx, row in enumerate(data_t1):
        for c_idx, val in enumerate(row):
            t1.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t1, [Inches(2.2), Inches(4.8)])

    add_h2("1.2. Vì sao anh Minh cần website")
    add_bullet("Mỗi ngày đội ngũ phải trả lời lặp đi lặp lại hàng chục tin nhắn hỏi cùng một vấn đề: đơn giá thiết kế bao nhiêu, có xưởng sản xuất trực tiếp không, chi phí thi công trọn gói tính thế nào, tiến độ làm mất bao lâu.")
    add_bullet("Gửi hồ sơ năng lực file PDF qua Zalo dung lượng rất nặng, khách xem trên điện thoại khó nhìn, không thể hiện được tính tương tác sinh động và độ sắc nét của các công trình thực tế đã bàn giao.")
    add_bullet("Khách hàng phân khúc cao cấp (chủ biệt thự biển, penthouse) đòi hỏi sự tin cậy tuyệt đối; việc thiếu một website chỉn chu khiến doanh nghiệp trông như một đội thợ nhận việc nhỏ lẻ, khó nhận được các hợp đồng lớn giá trị từ vài trăm triệu đến hàng tỷ đồng.")
    add_bullet("Các đơn vị thiết kế thi công đối thủ tại Nha Trang đã xây dựng website hiện đại, chuyên nghiệp nên anh lo ngại khách hàng tìm kiếm trên Google hay mạng xã hội sẽ chuyển sang đơn vị khác.")

    add_h2("1.3. Điều anh Minh lo ngại")
    add_bullet("Làm web xong nhìn đại trà, khô cứng, thiếu tính nghệ thuật (gu thẩm mỹ), không toát lên được phong cách đại dương hữu cơ độc bản của thương hiệu WHALE.")
    add_bullet("Web tải chậm, hình ảnh render và ảnh chụp thực tế công trình có độ phân giải cao bị vỡ nét hoặc làm đơ máy khi khách lướt trên điện thoại.")
    add_bullet("Mỗi lần cập nhật dự án mới bàn giao hoặc thay đổi bảng đơn giá vật liệu lại phải phụ thuộc vào lập trình viên và tốn thêm chi phí duy trì.")
    add_bullet("Chi phí đầu tư làm web cao mà không thu hút được khách hàng tiềm năng để lại thông tin đặt lịch tư vấn khảo sát.")

    add_h2("1.4. Điều anh Minh mong muốn")
    add_bullet("Một website đẳng cấp, mang ngôn ngữ tạp chí kiến trúc đương đại (Editorial Luxury), nhìn vào là thấy sự sang trọng, tinh tế và uy tín để khách hàng tin tưởng ngay từ 5 giây đầu tiên.")
    add_bullet("Làm nổi bật thế mạnh cạnh tranh cốt lõi: Làm chủ xưởng sản xuất Whale Craft 1.500m², giá gốc trực tiếp không qua trung gian, cam kết chất lượng thực tế giống 3D đến 98%.")
    add_bullet("Có tính năng tương tác độc đáo: Khám phá phòng tương tác thông minh (Room Explorer với điểm ghim sản phẩm) và Thước trượt so sánh 3D vs Bàn giao thực tế (Before/After Slider).")
    add_bullet("Tích hợp công cụ tính dự toán chi phí trực tuyến (Cost Calculator) giúp khách hàng tự ướm thử ngân sách theo diện tích và gói vật liệu trước khi liên hệ.")
    add_bullet("Khách hàng dễ dàng bấm gọi hotline, nhắn Zalo hoặc điền form đặt lịch khảo sát hiện trạng miễn phí.")
    add_bullet("Đội ngũ nhân viên văn phòng tự quản trị, đăng tải công trình mới và bài viết cẩm nang mà không cần biết kỹ thuật.")

    add_h2("1.5. Cách sử dụng phần giả định này")
    add_bullet("Các phần tiếp theo được xây dựng dựa trên chính những trăn trở và yêu cầu của anh Minh, đóng vai trò như bản yêu cầu chi tiết (Brief) gửi đến đội ngũ phát triển web.")
    add_bullet("Người thiết kế web căn cứ vào bản yêu cầu này để phân bổ sơ đồ trang web (Sitemap), xây dựng trải nghiệm người dùng (UX) và dựng cấu trúc giao diện hoàn chỉnh cho trang chủ.")

    # --- SECTION 2 ---
    add_h1("2. Tổng quan và bối cảnh")
    add_bullet("WHALE INTERIOR là thương hiệu thiết kế kiến trúc và thi công nội thất cao cấp tại Nha Trang – Khánh Hòa, chuyên sâu các dòng công trình: Biệt thự ven biển, Căn hộ Penthouse/Duplex, Nhà phố hiện đại và Khách sạn Boutique nghỉ dưỡng.")
    add_bullet("Thương hiệu sở hữu xưởng chế tác nội thất Whale Craft quy mô 1.500m² tại cụm công nghiệp Diên Phú với hệ thống máy cắt CNC 5 trục của Đức, phòng sơn áp suất dương và công nghệ sấy gỗ tự nhiên đạt chuẩn ẩm 8–12%. Các gói dịch vụ dao động từ 2.500.000đ/m² (Căn hộ tiêu chuẩn) đến 5.800.000đ/m² (Penthouse / Villa biển thượng hạng).")
    add_bullet("Hiện nay khách hàng chủ yếu biết đến công ty qua người quen giới thiệu hoặc Fanpage, gây quá tải trong việc tư vấn báo giá thủ công và hạn chế mở rộng tệp khách hàng tinh hoa. Chúng tôi cần một website thương hiệu chuẩn mực quốc tế để khẳng định năng lực, minh bạch quy trình, trình diễn công trình thực tế và thu thập khách hàng tiềm năng tự động 24/7.")

    add_h2("2.1. Tóm tắt yêu cầu")
    add_bullet("Thông điệp chính: “Vật liệu. Tĩnh tại. Không gian.” — Kiến trúc đương đại mang hơi thở đại dương sâu, chế tác chuẩn xác tại xưởng, giá trị trường tồn cùng thời gian.")
    add_bullet("Hành động ưu tiên: Khách hàng sử dụng công cụ Dự toán chi phí trực tuyến, trải nghiệm phòng mẫu Room Explorer và gửi yêu cầu Đặt lịch khảo sát hiện trạng miễn phí.")
    add_bullet("Định vị hình ảnh: Sang trọng, tối giản hữu cơ (Organic Minimalist), đĩnh đạc và đáng tin cậy; tông màu xanh biển sâu (Deep Ocean Navy #0A192F), màu cát ấm (Sand Cream #FAF8F5) và ánh kim đồng thau (Champagne Gold #C5A059).")
    add_bullet("Thiết bị ưu tiên: Điện thoại di động (Mobile First), tối ưu hóa trải nghiệm vuốt chạm mượt mà trên smartphone, song song với giao diện toàn cảnh chuẩn Cinematic trên Desktop.")

    # --- SECTION 3 ---
    add_h1("3. Khách hàng mục tiêu")
    add_bullet("Chúng tôi muốn website phục vụ bốn nhóm khách hàng trọng tâm. Mỗi nhóm có quy mô công trình, kỳ vọng thẩm mỹ và ngân sách khác nhau; do đó trang chủ phải trả lời nhanh chóng ba câu hỏi cốt tử: Doanh nghiệp có uy tín không, phong cách có hợp với tôi không, và tổng mức kinh phí ước tính là bao nhiêu.")

    t3 = doc.add_table(rows=5, cols=4)
    data_t3 = [
        ["Nhóm khách", "Đặc điểm", "Nhu cầu", "Website cần đáp ứng"],
        [
            "Chủ sở hữu Biệt thự biển, Villa nghỉ dưỡng",
            "35–60 tuổi; doanh nhân, Việt kiều, chủ doanh nghiệp; ngân sách từ 1.5 đến trên 5 tỷ đồng",
            "Đòi hỏi thiết kế độc bản, vật liệu cao cấp chịu được khí hậu biển (chống muối mặn, độ ẩm cao), tiến độ chuẩn xác",
            "Showroom ảnh dự án Biệt thự biển toàn cảnh, thông số vật liệu gỗ tự nhiên cao cấp, cam kết bảo hành 5 năm rõ ràng"
        ],
        [
            "Cư dân Penthouse, Căn hộ cao cấp",
            "28–45 tuổi; chuyên gia, quản lý cấp cao; phong cách sống hiện đại; ngân sách từ 300 đến 900 triệu",
            "Tối ưu hóa không gian mở, thẩm mỹ tinh tế chuẩn gu tạp chí quốc tế, đồ may đo chuẩn từng milimet",
            "Trải nghiệm Room Explorer tương tác từng món đồ, bộ lọc phong cách Modern / Japandi, bảng dự toán chi phí online"
        ],
        [
            "Gia chủ Nhà phố, Shophouse",
            "30–50 tuổi; gia đình đa thế hệ hoặc kết hợp kinh doanh; ngân sách 500 triệu đến 1.5 tỷ",
            "Bố trí công năng thông minh, vật liệu bền bỉ, giá xưởng minh bạch không bị đội giá trung gian",
            "Làm nổi bật lợi thế xưởng Whale Craft tiết kiệm 25–35%, quy trình 4 bước rõ ràng, hợp đồng phạt tiến độ"
        ],
        [
            "Chủ đầu tư Khách sạn, Homestay, Resort",
            "32–55 tuổi; đầu tư kinh doanh lưu trú ven biển Nha Trang, Cam Ranh, Quy Nhơn",
            "Năng lực thi công khối lượng lớn, tính đồng bộ cao, hoàn thiện sắc nét để đưa vào khai thác du lịch đúng hẹn",
            "Mục năng lực sản xuất xưởng 1.500m², máy móc CNC Đức, tiến độ cam kết và hotline kết nối trực tiếp KTS trưởng"
        ]
    ]
    for r_idx, row in enumerate(data_t3):
        for c_idx, val in enumerate(row):
            t3.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t3, [Inches(1.6), Inches(1.8), Inches(1.8), Inches(1.8)])

    add_h2("3.1. Hành vi chung của khách hàng")
    add_bullet("Khách hàng phân khúc trung và cao cấp luôn tìm hiểu rất kỹ lưỡng trước khi liên hệ; họ so sánh hồ sơ năng lực của ít nhất 3–4 đơn vị thiết kế thi công.")
    add_bullet("Đa số tiếp cận qua thiết bị di động khi đang tìm kiếm ý tưởng trên mạng hoặc được người quen gửi link website qua Zalo/Facebook.")
    add_bullet("Tâm lý e ngại: Sợ 'bản vẽ 3D thì lung linh nhưng thi công thực tế thì xấu', sợ bị phát sinh chi phí ngoài hợp đồng và sợ công trình bị chậm tiến độ bàn giao.")
    add_bullet("Cần sự minh bạch ngay lập tức: Muốn thấy hình ảnh công trình thực tế, thấy địa chỉ showroom và xưởng thật, có thể tính toán dự trù kinh phí sơ bộ mà chưa cần phải gọi điện thoại hỏi giá.")

    # --- SECTION 4 ---
    add_h1("4. Mục tiêu website")
    add_h2("4.1. Mục tiêu kinh doanh")
    add_bullet("Gia tăng 35–50% lượng khách hàng tiềm năng gửi yêu cầu tư vấn và khảo sát hiện trạng công trình hàng tháng.")
    add_bullet("Nâng cao giá trị hợp đồng bình quân thông qua việc định vị thương hiệu cao cấp, thu hút các công trình Villa và Penthouse quy mô lớn.")
    add_bullet("Cắt giảm 60% thời gian nhân sự phải giải thích các thông tin cơ bản (quy trình, năng lực xưởng, đơn giá m², chính sách bảo hành).")
    add_bullet("Khẳng định vị thế thương hiệu kiến trúc nội thất hữu cơ hàng đầu tại khu vực Nam Trung Bộ, cạnh tranh sòng phẳng với các đơn vị lâu năm.")

    add_h2("4.2. Mục tiêu người dùng")
    add_bullet("Cảm nhận được đẳng cấp thẩm mỹ và sự uy tín của thương hiệu ngay trong 5 giây đầu tiên truy cập.")
    add_bullet("Dễ dàng tra cứu các công trình thực tế theo đúng phong cách và loại hình nhà ở của mình.")
    add_bullet("Tự tính toán được mức chi phí dự kiến cho căn nhà của mình chỉ sau 3 bước chọn thông số đơn giản.")
    add_bullet("Đặt lịch hẹn gặp Kiến trúc sư trưởng hoặc liên hệ tư vấn qua Zalo/Hotline chỉ với 1 lần bấm duy nhất.")

    add_h2("4.3. Chỉ số đánh giá hiệu quả")
    t4 = doc.add_table(rows=5, cols=3)
    data_t4 = [
        ["Chỉ số", "Cách đo", "Mục tiêu"],
        ["Tỷ lệ bấm “Đặt lịch hẹn” / “Gọi hotline”", "Số lượt bấm trên tổng lượt truy cập website", "Đạt từ 6.5% – 10%"],
        ["Lượt sử dụng công cụ “Dự toán chi phí”", "Số lần người dùng thực hiện tính toán trên web", "Đạt trên 25% tổng lượt truy cập"],
        ["Số yêu cầu khảo sát hiện trạng gửi về", "Số form đăng ký tư vấn hoàn tất gửi về hệ thống", "Tăng trưởng đều 20–30% mỗi tháng"],
        ["Thời gian trung bình trên trang (Time on Site)", "Đo lường qua Google Analytics (GA4)", "Đạt trên 2 phút 30 giây (chứng tỏ khách xem kỹ dự án)"]
    ]
    for r_idx, row in enumerate(data_t4):
        for c_idx, val in enumerate(row):
            t4.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t4, [Inches(2.5), Inches(2.5), Inches(2.0)])

    # --- SECTION 5 ---
    add_h1("5. Ý tưởng chủ đạo và yêu cầu chức năng")
    add_h2("5.1. Ý tưởng chủ đạo")
    add_bullet("“Mắt thấy đẳng cấp – Tay chạm thực tế – Dự toán minh bạch”. Trang chủ dẫn dắt khách hàng qua một hành trình cảm xúc và lý trí chặt chẽ:")
    add_bullet("Thu hút thị giác bằng Hero Cinematic & Typography nghệ thuật → Tạo dựng lòng tin bằng số liệu năng lực → Khám phá không gian tương tác Room Explorer chạm vào từng sản phẩm → Chứng minh tay nghề chế tác bằng thước trượt Before/After (3D vs Thực tế) → Thuyết phục bằng xưởng sản xuất 1.500m² & quy trình 4 bước → Trao quyền cho khách tự tính dự toán ngân sách → Củng cố niềm tin bằng cẩm nang chuyên sâu → Thúc đẩy hành động Đặt lịch hẹn khảo sát miễn phí.")

    add_h2("5.2. Chức năng chúng tôi cần")
    add_bullet("Room Explorer tương tác điểm ghim (Hotspots): Khung cảnh phòng khách biệt thự với các điểm ghim phát xung nhịp; rê chuột/chạm vào điểm ghim sẽ hiện popup kính mờ thông tin sản phẩm, chất liệu, đơn giá và nút thêm vào giỏ hàng.")
    add_bullet("Thước trượt so sánh 3D và Thực tế (Before & After Slider): Kéo thanh trượt ngang để so sánh trực quan giữa bản vẽ thiết kế 3D và hình ảnh công trình thực tế hoàn thiện bàn giao.")
    add_bullet("Công cụ tính dự toán chi phí trực tuyến (Cost Calculator): Khách hàng nhập diện tích (m²), chọn loại hình công trình và gói vật liệu; hệ thống tự động tính toán tổng số tiền đầu tư ước tính theo thời gian thực.")
    add_bullet("Băng chuyền dự án dạng tạp chí kiến trúc (Editorial Carousel): Trình chiếu slide công trình tiêu biểu hiển thị kèm bảng mẫu vật liệu đặc trưng (Material Swatches) và bộ điều hướng mũi tên tròn.")
    add_bullet("Dải chữ chuyển động vô tận (Infinite Marquee Ticker): Chạy ngang các thông điệp khẳng định uy tín thương hiệu một cách tinh tế.")
    add_bullet("Trình phát âm thanh tĩnh lặng đại dương (Whale Soundscape): Nút bấm nổi bật/tắt âm thanh sóng biển êm dịu tạo cảm giác thư thái khi thưởng lãm không gian sống.")
    add_bullet("Giỏ hàng mô phỏng (Mini Cart Drawer): Cho phép khách hàng thêm đồ nội thất thiết kế từ phòng mẫu hoặc cửa hàng vào giỏ, xem trước tổng chi phí may đo.")
    add_bullet("Form đặt lịch khảo sát & tư vấn nhanh: Popup modal tinh gọn, xác thực số điện thoại và thông báo phản hồi KTS gọi lại trong 15 phút.")
    add_bullet("Hệ thống quản trị nội dung dễ dùng: Nhân viên công ty có thể tự đăng ảnh dự án mới, cập nhật bảng đơn giá và viết bài cẩm nang mà không cần can thiệp mã nguồn.")

    # --- SECTION 6 ---
    add_h1("6. Sitemap")
    add_bullet("Website gồm 9 nhóm trang chính, cấu trúc phẳng và mạch lạc để khách hàng chỉ mất tối đa 1 đến 2 cú nhấp chuột là tiếp cận được thông tin cần thiết.")

    t6 = doc.add_table(rows=10, cols=3)
    data_t6 = [
        ["Trang", "Nội dung và trang con", "Mục đích"],
        ["Trang chủ", "Tóm lược toàn bộ năng lực, trải nghiệm tương tác, dự án tiêu biểu, công cụ dự toán (xem mục 7)", "Định vị thương hiệu, khơi gợi cảm xúc và dẫn dắt khách đặt lịch tư vấn"],
        ["Giới thiệu & Xưởng", "Câu chuyện thương hiệu, triết lý hữu cơ, khám phá Xưởng Whale Craft 1.500m², máy CNC 5 trục, 4 cam kết vàng", "Xây dựng niềm tin vững chắc về năng lực sản xuất trực tiếp và chất lượng gốc"],
        ["Dự án / Công trình", "Danh mục dự án thực tế phân loại theo: Biệt thự biển, Penthouse, Nhà phố, Tân cổ điển Luxury kèm thông số diện tích và vật liệu", "Chứng minh năng lực thực chiến đa dạng qua các công trình bàn giao sắc nét"],
        ["Cửa hàng / Shop", "Bộ sưu tập đồ gỗ thủ công may đo: Ghế Whale Armchair, Bàn trà đá Marble, Giường gỗ óc chó, Kệ tivi phay sóng nước, giỏ hàng", "Giới thiệu dòng nội thất chế tác độc bản tại xưởng, phục vụ bán lẻ và may đo"],
        ["Kiến thức & Cẩm nang", "6 bài viết đúc kết kinh nghiệm: 7 sai lầm thiết kế, dự toán chi phí 2026, chọn gỗ cho vùng biển, kỹ thuật ánh sáng layering", "Chia sẻ giá trị, nâng cao độ tin cậy chuyên gia và tối ưu hóa SEO Google"],
        ["Dự toán chi phí", "Công cụ tương tác tính toán chi phí theo diện tích sàn, loại công trình và cấp độ vật liệu", "Minh bạch tài chính, giúp khách hàng chủ động ngân sách trước khi ký hợp đồng"],
        ["Đặt lịch tư vấn", "Form đăng ký thông tin khảo sát mặt bằng hiện trạng tại công trình, chọn thời gian phù hợp", "Thu thập thông tin khách hàng tiềm năng để đội ngũ KTS liên hệ chăm sóc"],
        ["Hỏi đáp & Chính sách", "Các câu hỏi thường gặp (FAQ), điều khoản hợp đồng thi công, cam kết bảo hành 5 năm & bảo trì trọn đời", "Giải tỏa triệt để các khúc mắc, đảm bảo tính minh bạch và pháp lý"],
        ["Liên hệ & Chi nhánh", "Trụ sở Showroom Nha Trang, Xưởng sản xuất Diên Khánh, Chi nhánh TP.HCM, Google Maps chỉ đường, Hotline 24/7", "Cung cấp địa chỉ thực tế rõ ràng để khách hàng ghé thăm showroom và xưởng"]
    ]
    for r_idx, row in enumerate(data_t6):
        for c_idx, val in enumerate(row):
            t6.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t6, [Inches(1.8), Inches(3.2), Inches(2.0)])

    # --- SECTION 7 ---
    add_h1("7. Cấu trúc trang chủ")
    add_bullet("Chúng tôi mong muốn trang chủ gồm 14 phân đoạn được sắp xếp tuần tự từ trên xuống dưới theo đúng quy luật tâm lý ra quyết định của khách hàng cao cấp:")

    t7 = doc.add_table(rows=15, cols=3)
    data_t7 = [
        ["TT", "Phần", "Nội dung chính và mục đích"],
        ["1", "Thanh thông báo trên cùng (Topbar)", "Dòng thông báo xưởng sản xuất trực tiếp 1.500m², số hotline 24/7 và email liên hệ chính thức."],
        ["2", "Đầu trang (Header)", "Logo tối giản biểu tượng đuôi cá voi, hệ thống menu điều hướng trái/phải, nút mở giỏ hàng và nút 'ĐẶT LỊCH HẸN'; tự động co giãn và làm mờ kính khi cuộn."],
        ["3", "Ảnh quảng cáo lớn (Hero Cinematic)", "Slideshow ảnh đại cảnh biệt thự tự động chuyển mượt mà, tiêu đề tạp chí 'VẬT LIỆU. Tĩnh Tại. KHÔNG GIAN.', nút kêu gọi trải nghiệm phòng tương tác và nhận báo giá."],
        ["4", "Đường lượn sóng & Chỉ số năng lực", "Dải đồ họa sóng biển chuyển tiếp mềm mại kết hợp 4 chỉ số tự động đếm số: 12+ Năm, 350+ Công trình, 1.500m² Xưởng, 100% Cam kết tiến độ."],
        ["5", "Dải chữ thương hiệu (Marquee Ticker)", "Dải chữ chạy ngang vô tận khẳng định tuyên ngôn: WHALE INTERIOR ✦ ORGANIC ARCHITECTURE ✦ XƯỞNG CHẾ TÁC 1.500M² ✦ CHUẨN XÁC 98% BẢN VẼ."],
        ["6", "Phòng mẫu tương tác (Room Explorer)", "Khung hình phòng khách biệt thự với 5 điểm ghim phát sáng (Hotspots) trên từng món nội thất; click xem popup giá xưởng, chất liệu và nút thêm vào giỏ."],
        ["7", "Bộ sưu tập tuyển chọn (Editorial Carousel)", "Slide công trình cao cấp dạng tạp chí kiến trúc với thông số 01/04, bảng mẫu vật liệu đặc trưng (Gỗ óc chó FAS, đá Calacatta) và nút chuyển slide."],
        ["8", "Dịch vụ toàn diện (Dark Wood Matrix)", "Khối nền màu gỗ óc chó sẫm sang trọng với 4 dịch vụ cốt lõi đánh số 01-04: Thiết kế kiến trúc, Xưởng Whale Craft, Thi công trọn gói và Bài trí nghệ thuật."],
        ["9", "So sánh 3D vs Thực tế (Before & After Slider)", "Khối so sánh kéo trượt ngang trực quan chứng minh thực tế bàn giao sắc nét giống 3D đến 98% tại Villa Ocean Horizon."],
        ["10", "Triết lý sáng tạo & Cam kết xưởng", "Khối giới thiệu giá trị cốt lõi và 4 ưu thế vượt trội: Tự sản xuất tại xưởng, bản vẽ chi tiết 1:1, tiết kiệm 25-35% chi phí, bảo hành dài hạn 5 năm."],
        ["11", "Công cụ dự toán trực tuyến (Cost Calculator)", "Hộp tính toán ngân sách tự động theo diện tích m², loại hình không gian (Chung cư/Nhà phố/Villa) và gói vật tư, hiển thị số tiền dự kiến tức thì."],
        ["12", "Cẩm nang kiến thức (Journal / Blog)", "Hiển thị 3 bài viết chuyên sâu đắt giá: 7 sai lầm thiết kế nội thất, dự toán chi phí xây nhà 2026, xu hướng Organic Coastal."],
        ["13", "Kêu gọi hành động lớn (CTA Banner)", "Biểu ngữ toàn màn hình thôi thúc khách hàng: 'Sẵn sàng tạo nên không gian vượt thời gian?' với nút đặt lịch tư vấn và gọi hotline."],
        ["14", "Chân trang & Nút liên hệ nổi (Footer & Floating)", "Thông tin showroom, địa chỉ xưởng, bản đồ chỉ đường, các nút liên hệ nổi cố định ở góc dưới: Âm thanh sóng biển (Whale Soundscape), Máy tính dự toán, Zalo và Hotline."]
    ]
    for r_idx, row in enumerate(data_t7):
        for c_idx, val in enumerate(row):
            t7.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t7, [Inches(0.6), Inches(2.3), Inches(4.1)])

    # --- SECTION 8 ---
    add_h1("8. Mong muốn về giao diện, Typography và trải nghiệm người dùng")
    add_bullet("Chúng tôi là những người làm nghề kiến trúc và chế tác mộc mỹ nghệ, không am hiểu sâu về công nghệ lập trình, nên xin bày tỏ cụ thể những kỳ vọng thẩm mỹ khi khách hàng ghé thăm ngôi nhà số của Whale Interior.")

    add_h2("8.1. Khách nhìn vào phải thấy gì")
    t8 = doc.add_table(rows=7, cols=2)
    data_t8 = [
        ["Điều chúng tôi quan tâm", "Mong muốn của chúng tôi"],
        ["Ấn tượng đầu tiên", "Vừa mở website là cảm nhận được ngay vẻ đẹp tĩnh tại, phóng khoáng của đại dương và chiều sâu của kiến trúc cao cấp; thấy rõ ngay năng lực làm chủ xưởng mộc 1.500m² mà không cần cuộn tìm lòng vòng."],
        ["Hình ảnh công trình", "Ảnh chụp kích thước lớn, sắc nét chuẩn Full HD/Retina, màu sắc chân thực. Tuyệt đối dùng ảnh công trình thật đã thi công của xưởng để khách đến tham quan thực tế không bị hụt hẫng."],
        ["Màu sắc & Chất cảm", "Bảng màu sang trọng: Deep Ocean Navy trầm ấm kết hợp Sand Cream dịu nhẹ và điểm nhấn Champagne Gold tinh tế; gợi cảm giác thư thái của biển khơi và sự ấm áp của chất liệu mộc."],
        ["Chữ và cách sắp xếp", "Bố cục tạp chí kiến trúc (Editorial Typography), kết hợp tinh tế giữa font Serif có chân đĩnh đạc và chữ sans-serif hiện đại. Cỡ chữ to rõ, khoảng cách dòng thoáng đãng, dễ đọc trên mọi thiết bị."],
        ["Sự tin cậy", "Địa chỉ Showroom và địa chỉ Xưởng tại Diên Khánh rõ ràng, định vị Google Maps chính xác, cam kết bảo hành 5 năm bằng hợp đồng pháp lý, bảng đơn giá minh bạch từng mét vuông."],
        ["Cá tính thương hiệu", "Phóng khoáng, tự do như cá voi đại dương nhưng tỉ mỉ, chuẩn xác và bền vững trong từng đường nét mộc. Sang trọng tinh tế nhưng không xa cách."]
    ]
    for r_idx, row in enumerate(data_t8):
        for c_idx, val in enumerate(row):
            t8.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t8, [Inches(2.3), Inches(4.7)])

    add_h2("8.2. Hệ thống Typography chuẩn mực kiến trúc cao cấp (Luxury Editorial Typography)")
    add_bullet("Hệ thống phông chữ được định chuẩn nghiêm ngặt theo tiêu chuẩn tạp chí kiến trúc và nhận diện thương hiệu WHALE, đảm bảo tính thẩm mỹ, sang trọng, tương phản cao và dễ đọc tuyệt đối trên mọi nền giao diện:")

    t8_typo = doc.add_table(rows=6, cols=3)
    data_t8_typo = [
        ["Thành phần chữ", "Quy cách font & Kích thước (Desktop / Tablet / Mobile)", "Yêu cầu thẩm mỹ & Khoảng cách (Hierarchy)"],
        ["Tiêu đề chính (Headlines - H1, H2)", "Font Serif cao cấp (Playfair Display), độ dày SemiBold (600).\n• Desktop: 58–72px (Hero) / 42–54px (Section)\n• Tablet: 46–58px / 34–44px\n• Mobile: 36–44px / 28–36px\nLine-height: 0.95–1.1; Letter-spacing: 0.01em", "Tăng độ tương phản tối đa với nền, không để chìm hoặc quá tối. Trên nền tối (Dịch vụ toàn diện, Dự toán) dùng màu Trắng tinh khiết (#FFFFFF) đổ bóng sâu. Hỗ trợ 100% tiếng Việt, không lỗi kerning hay rớt dòng."],
        ["Phụ đề / Eyebrow Text (ví dụ: 'DỊCH VỤ TOÀN DIỆN', 'DỰ TOÁN TRỰC TUYẾN')", "Font Sans-serif hiện đại (Plus Jakarta Sans), in hoa (Uppercase), độ dày SemiBold (600).\n• Cỡ chữ: 13–15px (chuẩn 14px)\n• Letter-spacing: 3–5px (chuẩn 4px)", "Hai đường line thanh mảnh (1px, 20–32px) nằm thẳng hàng và cân đối hai bên chữ tạo cảm giác chỉn chu, quý phái. Khoảng cách tới tiêu đề chính từ 25–35px (chuẩn 28px)."],
        ["Đoạn mô tả (Descriptions)", "Font Sans-serif hiện đại (Plus Jakarta Sans), độ dày Regular (400–500).\n• Desktop: 20–22px (Line-height: 1.6–1.8)\n• Tablet: 18px\n• Mobile: 16–18px (Line-height: 1.6)", "Màu chữ có độ tương phản cao (#334155 trên nền sáng, #E2E8F0 trên nền tối). Giới hạn chiều rộng 680–760px để người đọc không bị mỏi mắt. Cách tiêu đề 24–32px, cách nút CTA 32–40px."],
        ["Nút hành động (CTA Buttons)", "Font Sans-serif hiện đại (Plus Jakarta Sans), in hoa, độ dày SemiBold (600).\n• Cỡ chữ: 14–16px (chuẩn 15px)\n• Letter-spacing: 0.5–1px (chuẩn 0.8px)", "Chữ rõ nét và mạnh mẽ, căn giữa tuyệt đối cả hai chiều, icon mũi tên cách chữ 12px hợp lý. Hiệu ứng đổ bóng ánh kim Champagne Gold tinh tế, tạo cảm giác thủ công cao cấp."],
        ["Nhịp điệu phân cấp (Visual Hierarchy)", "Áp dụng đồng bộ tỉ lệ vàng thị giác cho toàn bộ 14 phân đoạn website", "Eyebrow ──[25–35px]──> Headline ──[24–32px]──> Description ──[32–40px]──> CTA Button. Tạo nhịp thở thị giác thoáng đãng, sang trọng."]
    ]
    for r_idx, row in enumerate(data_t8_typo):
        for c_idx, val in enumerate(row):
            t8_typo.rows[r_idx].cells[c_idx].paragraphs[0].text = val
    style_table(t8_typo, [Inches(2.1), Inches(2.6), Inches(2.3)])

    add_h2("8.3. Cơ chế Responsive tự nhiên 100% (Thuần Responsive – Không nút bấm thủ công)")
    add_bullet("Loại bỏ hoàn toàn các thanh công cụ chuyển đổi hay nút bấm giả lập thủ công: Website phải là một chỉnh thể Responsive tự nhiên 100%. Khi người dùng bấm phím F12 (Inspect / DevTools) trên trình duyệt máy tính, thu nhỏ cửa sổ hoặc truy cập trực tiếp bằng điện thoại di động thông minh, website tự động phát hiện và chuyển đổi mượt mà sang giao diện Mobile App.")
    add_bullet("Trải nghiệm Mobile App tối ưu trên thiết bị cầm tay: Tự động kích hoạt thanh điều hướng ứng dụng cố định cạnh dưới màn hình (Mobile App Bottom Navigation Bar) gồm: Trang chủ, Phòng mẫu tương tác, Dự án, Cửa hàng và Đặt lịch tư vấn.")
    add_bullet("Bảo toàn bố cục và phông chữ: Trên màn hình di động, phông chữ tự động điều chỉnh theo tỉ lệ chuẩn (Tiêu đề 36–44px, mô tả 16–18px), các khối lưới tự động xếp dọc (1 cột), bộ lọc chuyển sang dạng vuốt ngang (touch scroll), và thẻ popup sản phẩm Room Explorer tự động ghim đáy màn hình giúp thao tác bằng một ngón tay cái hoàn toàn tự nhiên và thuận tiện.")

    add_h2("8.4. Khách dùng phải thấy dễ & Hiệu năng vận hành")
    add_bullet("Tương tác mượt mà không độ trễ: Chạm vào điểm ghim Room Explorer là hiện ngay thẻ sản phẩm; kéo thanh trượt Before/After trơn tru như đang trải nghiệm ứng dụng cao cấp.")
    add_bullet("Kết nối nhanh chóng ở mọi vị trí: Dù đang ở bất kỳ trang nào hay cuộn đến phân đoạn nào, khách hàng luôn thấy nút gọi Hotline 24/7 và nút nhắn tin Zalo ở góc màn hình.")
    add_bullet("Tốc độ tải trang siêu tốc: Tối ưu dung lượng hình ảnh công trình và sử dụng hiệu ứng code thuần (CSS/JS) nhẹ nhàng để mở web trong vòng dưới 2 giây.")
    add_bullet("Quản trị trực quan cho nội bộ: Bàn giao giao diện quản trị bài viết và dự án đơn giản, nhân viên chỉ cần tải ảnh và gõ chữ là hiển thị đẹp chuẩn chỉnh lên website.")

    add_h2("8.5. Điều chúng tôi mong nhận được từ đơn vị thiết kế web")
    add_bullet("Bản demo trang chủ tương tác trực tiếp chạy thử trên môi trường web để ban giám đốc và đội ngũ kiến trúc sư duyệt trải nghiệm thực tế.")
    add_bullet("Sự hỗ trợ tinh chỉnh, lắng nghe phản hồi chi tiết về các góc nhìn hình ảnh và câu chữ thương hiệu sau khi xem bản mẫu.")
    add_bullet("Bộ tài liệu hướng dẫn và bàn giao mã nguồn rõ ràng, giúp đội ngũ Whale Interior hoàn toàn chủ động làm chủ website về lâu dài.")

    output_path = r"C:\Users\ngogi\.gemini\antigravity\scratch\whale-interior\Brief_Website_Whale_Interior.docx"
    try:
        doc.save(output_path)
        print(f"Document saved successfully to {output_path}")
    except PermissionError:
        output_alt = r"C:\Users\ngogi\.gemini\antigravity\scratch\whale-interior\Brief_Website_Whale_Interior_CapNhat.docx"
        doc.save(output_alt)
        print(f"File Brief_Website_Whale_Interior.docx dang duoc mo trong trinh doc (WPS/Word). Da luu thanh cong ban cap nhat tai: {output_alt}")

if __name__ == "__main__":
    create_whale_brief_doc()
